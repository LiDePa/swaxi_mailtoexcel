from openpyxl import load_workbook
import os
import sys
import re


#MITARBEITERLISTE:

class Mitarbeiter:
	def __init__(self, fahrernummer, vorname, nachname, mailadressen, kurzname):
		self.fahrernummer = fahrernummer
		self.vorname = vorname
		self.nachname = nachname
		self.mailadressen = mailadressen
		self.kurzname = kurzname
	def __repr__(self):
		return f"<Mitarbeiter {self.kurzname} ({self.fahrernummer})>"
	
# Hilfsfunktion für korrekten Pfad (auch für PyInstaller)
def get_resource_path(filename):
	if hasattr(sys, '_MEIPASS'):
		return os.path.join(sys._MEIPASS, filename)
	return os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)

# Mitarbeiterdaten laden und fehlende kurznamen ergänzen
def lade_mitarbeiter():
	MITARBEITER_DATEI = "Mitarbeiter.xlsx"
	MITARBEITER_PATH = get_resource_path(MITARBEITER_DATEI)
	wb = load_workbook(MITARBEITER_PATH)
	ws = wb.active
	mitarbeiter_liste = []
	for row in ws.iter_rows(min_row=2):  # Erste Zeile ist Header
		# Stoppe, wenn die Zeile komplett leer ist
		if all(cell.value is None for cell in row):
			break
		fahrernummer = row[0].value  # Spalte A: "Titel"
		vorname = row[1].value      # Spalte B: "Vorname"
		nachname = row[2].value     # Spalte C: "Nachname"
		mailadressen = row[5].value # Spalte F: "E-Mail Privat"
		kurzname = row[10].value    # Spalte K: "kurzname"
		# Validierung
		if not fahrernummer or (not vorname and not nachname):
			print(f"Fehlerhafte Zeile: Fahrernummer oder Name fehlt. Ignoriert: {vorname} {nachname}")
			continue
		# kurzname generieren falls leer
		if not kurzname:
			kurzname = generiere_kurzname(vorname, nachname)
			row[10].value = kurzname
		mitarbeiter = Mitarbeiter(fahrernummer, vorname, nachname, mailadressen, kurzname)
		mitarbeiter_liste.append(mitarbeiter)
	wb.save(MITARBEITER_PATH)
	return mitarbeiter_liste

# kurzname generieren
def generiere_kurzname(vorname, nachname):
	vornamen = str(vorname).split()
	nachnamen = str(nachname).split()
	initialen = [f"{v[0]}." for v in vornamen if v]
	kurzname = " ".join(initialen + nachnamen)
	return kurzname.strip()





# ANONYMISIERUNG:

def _zerlege_name(name):
    """Zerlegt einen Namen (z.B. 'Linus Lorenzo') in einzelne Bestandteile (z.B. ['Linus', 'Lorenzo'])."""
    if not name:
        return []
    # Trennen nach Leerzeichen oder Bindestrich
    return re.split(r'[\s\-]+', str(name).strip())

def erstelle_mitarbeiter_pattern(mitarbeiter_liste):
    """
    Build a single regex pattern for all employee name parts (first, last).
    Full names are deliberately excluded to reduce pattern size.
    """
    alle_namen = set()
    for m in mitarbeiter_liste:
        # Nur die einzelnen Namensbestandteile hinzufügen
        for teil in _zerlege_name(m.vorname):
            alle_namen.add(teil)
        for teil in _zerlege_name(m.nachname):
            alle_namen.add(teil)
    
    # Entferne leere Strings
    alle_namen.discard("")
    
    # Sortierung nach Länge ist hier nicht mehr kritisch, da wir nur Einzelnamen haben, 
    # aber wir behalten es bei, um maximale Namensteil-Länge zu priorisieren (z.B. Doppelnamen).
    sortierte_namen = sorted(list(alle_namen), key=len, reverse=True)
    
    # Füge \b (Wortgrenzen) um jeden Namen hinzu
    escaped_namen = [re.escape(name) for name in sortierte_namen]
    
    pattern_str = r"\b(" + "|".join(escaped_namen) + r")\b"
    return re.compile(pattern_str, flags=re.IGNORECASE)


def anonymisiere_mailtexte(mailjson_liste, mitarbeiter_liste):
    """
	Replace all employee names and anonymize all email addresses in the mailtext.
    """
    # 1. Mitarbeiter-Namen-Pattern erstellen
    mitarbeiter_pattern = erstelle_mitarbeiter_pattern(mitarbeiter_liste)
    
    # 2. Pattern für mailadressen (@ enthalten)
    mail_adresse_pattern = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"

    for mail in mailjson_liste:
        mail_text_bearbeitet = mail.mailtext or ""

        # Mitarbeiter-Namen anonymisieren
        mail_text_bearbeitet = mitarbeiter_pattern.sub("Anon", mail_text_bearbeitet)

        # mailadressen anonymisieren
        mail_text_bearbeitet = re.sub(mail_adresse_pattern, "mailadress_anon", mail_text_bearbeitet, flags=re.IGNORECASE)

        mail.mailtext_anonym = mail_text_bearbeitet.strip()

    return mailjson_liste