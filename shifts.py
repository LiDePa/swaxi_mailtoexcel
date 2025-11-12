"""Hilfsfunktionen für Schichtwunsch-Tabellen."""

from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import Workbook, load_workbook


_WOCHENTAGE: Tuple[str, ...] = (
	"Mittwoch",
	"Donnerstag",
	"Freitag",
	"Samstag",
	"Sonntag",
	"Montag",
	"Dienstag",
)


def _to_sequence(value) -> Sequence[str]:
	"""Normalisiert beliebige Schichtwerte zu einer Sequenz von Strings."""
	if value is None:
		return ()
	if isinstance(value, str):
		parts = [p.strip() for p in value.replace(";", ",").split(",")]
		clean = [p for part in parts for p in part.split() if p]
		return tuple(clean)
	if isinstance(value, dict):
		return tuple(str(k).strip() for k, v in value.items() if v)
	if isinstance(value, Iterable):
		return tuple(str(item).strip() for item in value if str(item).strip())
	return (str(value).strip(),)


def _sanitize_shift(shift: str) -> str:
	return shift.replace(".", "").strip()


def erzeuge_schicht_tabellen(mails) -> Tuple[List[List[str]], List[List[str]], List[Dict[str, str]]]:
	"""Erstellt Schicht- und Mail-ID-Tabellen sowie Kommentarlisten."""
	spalteninhalte = defaultdict(list)
	mail_id_spalten = defaultdict(list)
	kommentar_map: Dict[str, str] = {}
	kommentare_liste: List[Dict[str, str]] = []

	for mail in mails or []:
		if not getattr(mail, "ist_schichtwunsch", False):
			continue

		schichtwuensche = getattr(mail, "schichtwuensche", {}) or {}
		if not isinstance(schichtwuensche, dict):
			continue
		kurzname = str(getattr(mail, "kurzname", "")).strip()
		mail_id = getattr(mail, "mail_id", "")

		kommentar = str(getattr(mail, "kommentar", "") or "").strip()
		if kommentar:
			if kurzname:
				kommentar_map[kurzname] = kommentar
			kommentare_liste.append({"kurzname": kurzname, "kommentar": kommentar, "mail_id": str(mail_id)})

		for tag in _WOCHENTAGE:
			raw_shift_values = schichtwuensche.get(tag)
			shifts = []
			for item in _to_sequence(raw_shift_values):
				cleaned = _sanitize_shift(item)
				if cleaned:
					shifts.append(cleaned)
			if not shifts:
				continue

			eintrag = f"{kurzname} {','.join(shifts)}".strip()
			spalteninhalte[tag].append(eintrag)
			mail_id_spalten[tag].append(str(mail_id))

	max_zeilen = max((len(spalteninhalte[tag]) for tag in _WOCHENTAGE), default=0)

	shifts_table = [list(_WOCHENTAGE)]
	mail_ids_table = [list(_WOCHENTAGE)]

	for zeile in range(max_zeilen):
		shifts_row = []
		mail_ids_row = []
		for tag in _WOCHENTAGE:
			if zeile < len(spalteninhalte[tag]):
				shifts_row.append(spalteninhalte[tag][zeile])
				mail_ids_row.append(mail_id_spalten[tag][zeile])
			else:
				shifts_row.append("")
				mail_ids_row.append("")
		shifts_table.append(shifts_row)
		mail_ids_table.append(mail_ids_row)

	return shifts_table, mail_ids_table, kommentare_liste


def schreibe_schichten_nach_excel(
	shifts_table: List[List[str]],
	kommentare: List[Dict[str, str]],
	kalenderwoche: int,
) -> Path:
	"""Schreibt die Schichttabelle und Kommentare in die Excel-Datei.

	Args:
		shifts_table: Tabelle mit Kopfzeile und Schichtwerten.
		kommentare: Liste der Kommentare inklusive Kurznamen.
		kalenderwoche: Kalenderwoche, die als Blattname verwendet wird.

	Returns:
		Pfad zur aktualisierten Excel-Datei.

	Raises:
		ValueError: Wenn Kalenderwoche fehlt oder das Blatt bereits existiert.
	"""
	if not kalenderwoche:
		raise ValueError("Kalenderwoche fehlt für den Excel-Export.")

	sheet_name = f"KW{int(kalenderwoche):02d}"
	workbook_path = Path(__file__).with_name("Fahrer Dienstwünsche.xlsx")

	if workbook_path.exists():
		workbook = load_workbook(workbook_path)
		if sheet_name in workbook.sheetnames:
			raise ValueError(f"Blatt '{sheet_name}' existiert bereits.")
		worksheet = workbook.create_sheet(title=sheet_name)
	else:
		workbook = Workbook()
		worksheet = workbook.active
		worksheet.title = sheet_name

	# Erste Zeile bleibt bewusst leer, daher beginnt die Kopfzeile in Zeile 2.
	headers = shifts_table[0] if shifts_table else list(_WOCHENTAGE)
	for column_index in range(len(_WOCHENTAGE)):
		value = headers[column_index] if column_index < len(headers) else _WOCHENTAGE[column_index]
		worksheet.cell(row=2, column=column_index + 1, value=value)

	for row_offset, row_values in enumerate(shifts_table[1:], start=3):
		for column_index, cell_value in enumerate(row_values[: len(_WOCHENTAGE)]):
			worksheet.cell(row=row_offset, column=column_index + 1, value=cell_value)

	# Spalte H bleibt frei, Kommentare starten in I2.
	worksheet.cell(row=2, column=9, value="Kommentare")
	for comment_offset, entry in enumerate(kommentare or [], start=3):
		kurzname = (entry.get("kurzname") or "").strip()
		kommentar = (entry.get("kommentar") or "").strip()
		text = kommentar
		if kurzname and kommentar:
			text = f"{kurzname}: {kommentar}"
		elif kurzname:
			text = kurzname
		worksheet.cell(row=comment_offset, column=9, value=text)

	worksheet.freeze_panes = "A3"
	workbook.save(workbook_path)
	return workbook_path

